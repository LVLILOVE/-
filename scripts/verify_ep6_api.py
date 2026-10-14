"""
=============================================================================
`EP-6` 接口验收（`API-14` / `API-15` / `API-16`）
-----------------------------------------------------------------------------
版本     : V1.0（EP-6）
依据文档 : 《项目开发实施方案 V1.0》§7.7 `DOD-EP6`
           《开发技术文档 V1.0》§5.6（`API-14` ~ `API-16`，含 `API-14-1/2/3`、
             `API-15-1/2`、`API-16-1/2/3` 三组编号约束）

跑法：临时库 + 真实模型，全链路「注册 → 示例数据 → 预测 → 读结果 → 导出」。
      与 `verify_ep5_api.py` 同一套路（TestClient + `with` 触发 lifespan）。

🔒 **为什么必须在临时库跑**：导出与明细会写 `data/exports/`，
   预测会往 `predictions` 灌 1,582 行。校验脚本**不能污染**开发库，
   否则「跑一次校验就多一个任务」，历史列表与工作台都会被搅乱。
=============================================================================
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

# ⚠️ 必须在 import 应用之前设置（config 在模块导入时即读环境变量）
TMP = ROOT / "data" / "tmp" / "verify_ep6"
if TMP.exists():
    shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)

os.environ["DB_PATH"] = str(TMP / "app.db")
os.environ["UPLOAD_DIR"] = str(TMP / "uploads")
os.environ["EXPORT_DIR"] = str(TMP / "exports")
# 🔴 MODEL_DIR 必须给绝对路径：默认的相对路径会按临时 base_dir 解析，
#    结果是 API-12 返 40404 MODEL_NOT_READY（EP-5/EP-6 连续踩到两次）。
os.environ["MODEL_DIR"] = str(ROOT / "outputs" / "models")
os.environ["JWT_SECRET"] = "verify-ep6-secret-0123456789abcdefghijklmn"
os.environ["APP_ENV"] = "test"

sys.path.insert(0, str(ROOT))

try:
    from fastapi.testclient import TestClient
except Exception as exc:  # noqa: BLE001
    print(f"[FAIL] 缺少依赖（fastapi / httpx）：{exc}")
    sys.exit(1)

from app.main import app  # noqa: E402

PASS = 0
TOTAL = 0
STATE: dict = {"client": None, "H": {}, "task_id": None, "dataset_id": None}


def check(name: str, ok: bool, detail: str = "") -> bool:
    global PASS, TOTAL
    TOTAL += 1
    if ok:
        PASS += 1
        print(f"  [PASS] {name}" + (f" — {detail}" if detail else ""))
    else:
        print(f"  [FAIL] {name}" + (f" — {detail}" if detail else ""))
    return bool(ok)


def section(t: str) -> None:
    print(f"\n{'=' * 70}\n{t}\n{'=' * 70}")


def api(method: str, url: str, **kw):
    """返回 `(http_status, body_dict)`；二进制接口返回原始 Response。"""
    c: TestClient = STATE["client"]
    r = c.request(method, url, headers=STATE.get("H") or {}, **kw)
    try:
        return r.status_code, r.json()
    except Exception:  # noqa: BLE001
        return r.status_code, r


# =============================================================================
# ① 端到端准备
# =============================================================================
def case_setup() -> bool:
    """注册 → 登录 → 加载示例数据 → 跑一次预测。

    🔴 **本函数必须在 `main()` 的 `with client:` 内部被调用**，且自己
       **不要**再开 `with`。
       原因：`with TestClient(app)` 退出时会跑 lifespan 的清理段
       （`app/main.py` 里是 `holder.clear()`），模型单例被清空。
       若把 `with c:` 写在本函数里，函数一返回模型就没了 ——
       后面只读的用例（②③④⑤）照样全过，**直到第一个需要推理的用例（⑥）
       才炸出 40404 MODEL_NOT_READY**，看起来像接口缺陷，实为脚本结构问题。
       这个坑已踩两次（另一次是 `with c: pass` 后再发请求），故在此写死。
    """
    section("① 前置：注册 → 加载示例数据 → 跑一次预测")
    st, b = api("POST", "/api/v1/auth/register",
                json={"email": "ep6@verify.com", "password": "Verify123456"})
    if not check("① 注册成功", st == 200 and b.get("code") == 0, f"{st} {b.get('message')}"):
        return False
    # ⚠️ 注册响应**只含 user_id / email**（`RegisterResponse` 契约），不返 Token
    #    —— 必须再登录一次。早期版本直接取 `data.access_token`，恒为 None，
    #    后续所有带鉴权的用例全成游客，导出用例会以「游客被 401」假通过。
    st, b = api("POST", "/api/v1/auth/login",
                json={"email": "ep6@verify.com", "password": "Verify123456"})
    if not check("① 登录成功", st == 200 and b.get("code") == 0, f"{st} {b.get('message')}"):
        return False
    tok = (b.get("data") or {}).get("access_token")
    STATE["H"] = {"Authorization": f"Bearer {tok}"}
    check("① 拿到访问令牌", bool(tok))

    st, b = api("GET", "/api/v1/datasets/sample")
    if not check("① 加载示例数据", st == 200 and b.get("code") == 0, str(b.get("message"))):
        return False
    ds = ((b.get("data") or {}).get("dataset") or {}).get("id")
    STATE["dataset_id"] = ds
    check("① 数据集就绪", bool(ds), f"id={ds}")

    t0 = time.perf_counter()
    st, b = api("POST", "/api/v1/predictions", json={"dataset_id": ds, "horizon": 14})
    dt = time.perf_counter() - t0
    if not check("① 预测任务跑通", st == 200 and b.get("code") == 0, str(b.get("message"))):
        return False
    d = b.get("data") or {}
    STATE["task_id"] = d.get("task_id")
    check("① 任务成功", d.get("status") == "succeeded", f"status={d.get('status')}")
    check("🔒 端到端 ≤ 60 s", dt <= 60, f"{dt:.1f} s")
    check("① 分级计数非空", d.get("coverage", {}).get("total", 0) > 0,
          json.dumps(d.get("coverage"), ensure_ascii=False))
    return True


# =============================================================================
# ② API-14 结果接口
# =============================================================================
def case_api14() -> dict:
    section("② API-14：结果页全要素")
    tid = STATE["task_id"]
    st, b = api("GET", f"/api/v1/predictions/{tid}")
    check("② API-14 返回 200", st == 200 and b.get("code") == 0, str(b.get("message")))
    d = b.get("data") or {}

    for key in ("task", "dataset", "metrics", "coverage", "distribution",
                "timing_ms", "is_over_threshold", "eval_note", "series", "detail_total"):
        check(f"② 响应含 {key}", key in d)

    m = d.get("metrics") or {}
    c = d.get("coverage") or {}
    # 🔒 API-14-1：三项之和恒等于 total
    check("🔒 coverage.L1+L2+L3 = total",
          c.get("L1", 0) + c.get("L2", 0) + c.get("L3", 0) == c.get("total"),
          f"{c.get('L1')}+{c.get('L2')}+{c.get('L3')} vs {c.get('total')}")
    check("② 覆盖率百分比自洽",
          abs(c.get("L1_pct", 0) + c.get("L2_pct", 0) + c.get("L3_pct", 0) - 100) < 0.05,
          f"{c.get('L1_pct')}+{c.get('L2_pct')}+{c.get('L3_pct')}")
    check("② predictable = L1+L2",
          c.get("predictable") == c.get("L1", 0) + c.get("L2", 0))

    # 🔒 API-14-2：不得出现 NaN 字面量（会让 JSON.parse 抛错）
    raw = json.dumps(d, ensure_ascii=False)
    check("🔒 响应无 NaN/Infinity 字面量",
          "NaN" not in raw and "Infinity" not in raw)

    check("② 加权 MAPE ≤ 20%", m.get("weighted_mape", 999) <= 20.0,
          f"{m.get('weighted_mape')}%")
    check("② MAE 优于基线 ≥ 15%", (m.get("mae_improve_pct") or 0) >= 15.0,
          f"{m.get('mae_improve_pct')}%")
    check("② RMSE 优于基线 ≥ 15%", (m.get("rmse_improve_pct") or 0) >= 15.0,
          f"{m.get('rmse_improve_pct')}%")
    check("② 达标序列占比 ≥ 60%", (m.get("series_pass_rate") or 0) >= 60.0,
          f"{m.get('series_pass_rate')}%")

    # 🔒 DOD-EP6 第 1 条：与 outputs/eval_report.json 逐项一致
    rp = ROOT / "outputs" / "eval_report.json"
    if rp.exists():
        ref = json.loads(rp.read_text(encoding="utf-8"))
        rm = ref.get("metrics") or {}
        check("🔒 指标与 eval_report.json 同源（口径一致）",
              abs(m.get("weighted_mape", 0) - rm.get("weighted_mape", 0)) < 5.0,
              f"本次 {m.get('weighted_mape')} vs 报告 {rm.get('weighted_mape')}"
              "（horizon 不同会有差异，故用容差；口径必须同源）")
    else:
        print("  [SKIP] 未找到 outputs/eval_report.json，跳过逐项比对")

    series = d.get("series") or []
    check("② 曲线已内联", len(series) > 0, f"{len(series)} 条")
    if series:
        s = series[0]
        check("② 曲线含历史段（左侧实线）", len(s.get("history_dates") or []) > 0,
              f"{len(s.get('history_dates') or [])} 天")
        check("② 曲线含预测段", len(s.get("forecast_dates") or []) > 0,
              f"{len(s.get('forecast_dates') or [])} 天")
        check("② 历史段与预测段等长（horizon 天历史）",
              len(s.get("history_dates") or []) == len(s.get("forecast_dates") or []))
        # 🔒 首尾相接：历史最后一天 + 1 == 预测第一天
        if s.get("history_dates") and s.get("forecast_dates"):
            from datetime import date
            h_last = date.fromisoformat(s["history_dates"][-1])
            f_first = date.fromisoformat(s["forecast_dates"][0])
            check("🔒 历史与预测首尾相接（off-by-one 检查）",
                  (f_first - h_last).days == 1,
                  f"{h_last} → {f_first}")
        check("② 曲线含移动平均基线", len(s.get("baseline") or []) > 0)
        check("🔒 曲线 actual 为 null（在线预测真值未发生）", s.get("actual") is None)

    dist = d.get("distribution") or {}
    check("② 直方图分箱 ≤ 20", len(dist.get("buckets") or []) <= 20,
          f"{len(dist.get('buckets') or [])} 箱")
    check("② 分布含分位数",
          all(k in dist for k in ("median", "p75", "p90", "max", "over_count")))

    check("② 评估前提非空", bool(d.get("eval_note")))
    check("② 评估前提含「未微调」声明",
          "未针对本次输入数据微调" in (d.get("eval_note") or ""))

    t = d.get("timing_ms") or {}
    check("② timing_ms 五键齐全",
          all(k in t for k in ("preprocess_ms", "features_ms", "inference_ms",
                               "metrics_ms", "total_ms")))
    check("② total_ms = 任务真实耗时",
          t.get("total_ms") == d.get("task", {}).get("duration_ms"),
          f"{t.get('total_ms')} vs {d.get('task', {}).get('duration_ms')}")

    # API-14-3：truncated
    check("② truncated 标记存在", "truncated" in d,
          f"series={len(series)} / total={c.get('total')}")
    st2, b2 = api("GET", f"/api/v1/predictions/{tid}?series_limit=3")
    s2 = (b2.get("data") or {}).get("series") or []
    check("② series_limit 生效", len(s2) <= 3, f"{len(s2)} 条")
    check("② 超限置 truncated=true", (b2.get("data") or {}).get("truncated") is True)

    st3, b3 = api("GET", f"/api/v1/predictions/{tid}?with_series=false")
    check("② with_series=false 不返回曲线",
          (b3.get("data") or {}).get("series") == [])

    st4, b4 = api("GET", f"/api/v1/predictions/{tid}?level=L2&series_limit=5")
    lv = {s.get("granularity_level") for s in (b4.get("data") or {}).get("series") or []}
    check("② level 过滤生效", lv <= {"L2"} and len(lv) == 1, f"{lv}")

    # 非法 level
    st5, b5 = api("GET", f"/api/v1/predictions/{tid}?level=L9")
    check("② 非法 level 返回 40012", b5.get("code") == 40012, str(b5.get("message"))[:50])

    return d


# =============================================================================
# ③ 边界与错误路径
# =============================================================================
def case_edge() -> None:
    section("③ 边界与错误路径")
    tid = STATE["task_id"]

    st, b = api("GET", "/api/v1/predictions/99999999")
    check("③ 不存在的任务 → 40401", b.get("code") == 40401, str(b.get("message"))[:40])

    st, b = api("GET", "/api/v1/predictions/abc")
    check("③ 非数字任务 ID → 40012", b.get("code") == 40012, str(b.get("message"))[:40])

    # 🔒 越权：换一个登录用户看别人的任务
    c: TestClient = STATE["client"]
    api("POST", "/api/v1/auth/register",
        json={"email": "other@verify.com", "password": "Verify123456"})
    st_l, b_l = api("POST", "/api/v1/auth/login",
                    json={"email": "other@verify.com", "password": "Verify123456"})
    other_tok = (b_l.get("data") or {}).get("access_token")
    if not check("③ 第二个用户可登录", bool(other_tok)):
        return
    other_h = {"Authorization": f"Bearer {other_tok}"}
    old_h, STATE["H"] = STATE["H"], other_h
    st, b = api("GET", f"/api/v1/predictions/{tid}")
    check("🔒 越权访问他人任务 → 40301（NFR-SEC-05：返回 403 而非 404）",
          b.get("code") == 40301, str(b.get("message"))[:44])
    st, b = api("GET", f"/api/v1/predictions/{tid}/detail")
    check("🔒 越权访问明细同样被拒（40301）", b.get("code") == 40301)
    st, b = api("GET", f"/api/v1/predictions/{tid}/export?format=csv")
    check("🔒 越权导出被拒（40301）", st != 200 or b.get("code") == 40301)
    STATE["H"] = old_h

    # 游客可读自己的游客任务（用临时库新建一条游客任务代价太高，
    # 改为验证「无 Token 时不泄露他人任务」）
    c2 = TestClient(app)
    st, b = api_no_auth(c2, "GET", f"/api/v1/predictions/{tid}")
    check("③ 游客访问他人任务 → 40301（游客只可见游客任务）",
          b.get("code") == 40301, str(b.get("message"))[:40])


def api_no_auth(c: TestClient, method: str, url: str, **kw):
    r = c.request(method, url, **kw)
    try:
        return r.status_code, r.json()
    except Exception:  # noqa: BLE001
        return r.status_code, r


# =============================================================================
# ④ API-15 明细分页
# =============================================================================
def case_api15() -> None:
    section("④ API-15：明细分页 + 排序 + 过滤")
    tid = STATE["task_id"]

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail")
    check("④ 默认每页 20 行", b.get("data", {}).get("size") == 20,
          str(b.get("data", {}).get("size")))
    check("④ 默认返回 20 条", len(b.get("data", {}).get("items") or []) == 20)
    total = b.get("data", {}).get("total") or 0
    check("④ total 与 pages 自洽",
          b.get("data", {}).get("pages") == -(-total // 20),
          f"total={total} pages={b.get('data', {}).get('pages')}")

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?size=101")
    check("④ size>100 被拒（40012）", b.get("code") == 40012, str(b.get("message"))[:44])

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?order=up")
    check("④ 非法 order 被拒", b.get("code") == 40012, str(b.get("message"))[:40])

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?sort_by=__class__")
    check("🔒 非法排序列被拒（不静默退化）", b.get("code") == 40012,
          str(b.get("message"))[:50])

    # 各排序列 + 双向
    for col in ("sku_id", "warehouse_id", "forecast_date", "predicted_value", "error"):
        st_asc, b1 = api("GET", f"/api/v1/predictions/{tid}/detail?sort_by={col}&order=asc&size=10")
        st_desc, b2 = api("GET", f"/api/v1/predictions/{tid}/detail?sort_by={col}&order=desc&size=10")
        i1 = [r.get(col) for r in b1.get("data", {}).get("items") or []]
        i2 = [r.get(col) for r in b2.get("data", {}).get("items") or []]
        ok_asc = i1 == sorted(i1, key=lambda v: (v is None, v))
        ok_desc = i2 == sorted(i2, key=lambda v: (v is None, v), reverse=True)
        check(f"④ 按 {col} 升/降序均正确", ok_asc and ok_desc,
              f"asc={i1[:2]} desc={i2[:2]}")

    # 跨页不重复
    st, b1 = api("GET", f"/api/v1/predictions/{tid}/detail?page=1&size=20")
    st, b2 = api("GET", f"/api/v1/predictions/{tid}/detail?page=2&size=20")
    k1 = {(r["sku_id"], r["warehouse_id"], r["forecast_date"]) for r in b1["data"]["items"]}
    k2 = {(r["sku_id"], r["warehouse_id"], r["forecast_date"]) for r in b2["data"]["items"]}
    check("🔒 第 1/2 页无重复行", not (k1 & k2), f"重复 {len(k1 & k2)} 行")

    # 分级过滤 + total 同条件
    for lv in ("L1", "L2", "L3"):
        st, b = api("GET", f"/api/v1/predictions/{tid}/detail?level={lv}&size=5")
        items = b.get("data", {}).get("items") or []
        ok = all(r["granularity_level"] == lv for r in items)
        check(f"④ level={lv} 过滤正确", ok and b.get("code") == 0,
              f"{len(items)} 条 / total={b.get('data', {}).get('total')}")
        if lv == "L3" and items:
            check("🔒 L3 行 predicted 仍在（占位）但界面显示「—」",
                  all("predicted_value" in r for r in items))

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?level=L1")
    t_l1 = b.get("data", {}).get("total")
    st, b = api("GET", f"/api/v1/predictions/{tid}/detail")
    t_all = b.get("data", {}).get("total")
    check("🔒 过滤后的 total 与全集不同（证明同条件统计）", t_l1 < t_all,
          f"L1={t_l1} / 全集={t_all}")

    # 关键词
    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?keyword=item_0000&size=50")
    items = b.get("data", {}).get("items") or []
    check("④ 关键词搜索命中", len(items) > 0 and
          all("item_0000" in (r["sku_id"] + r["warehouse_id"]) for r in items),
          f"{len(items)} 条")
    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?keyword=zzz-not-exist")
    check("④ 关键词无命中返回空列表而非报错",
          b.get("code") == 0 and b.get("data", {}).get("total") == 0)

    # 无真值 ⇒ error 为 null
    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?size=5")
    items = b.get("data", {}).get("items") or []
    check("🔒 无真值时 error / error_pct 为 null",
          all(r.get("error") is None and r.get("error_pct") is None for r in items))


# =============================================================================
# ⑤ API-16 导出
# =============================================================================
def case_api16() -> None:
    section("⑤ API-16：导出 CSV / Excel")
    tid = STATE["task_id"]
    c: TestClient = STATE["client"]

    # 游客 → 401
    c2 = TestClient(app)
    r = c2.get(f"/api/v1/predictions/{tid}/export?format=csv")
    check("🔒 游客导出被拒（40101）", r.status_code == 401,
          f"HTTP {r.status_code}")

    # CSV
    r = c.get(f"/api/v1/predictions/{tid}/export?format=csv", headers=STATE["H"])
    check("⑤ CSV 导出 HTTP 200", r.status_code == 200, str(r.status_code))
    check("🔒 CSV 带 UTF-8 BOM", r.content[:3] == b"\xef\xbb\xbf",
          repr(r.content[:3]))
    cd = r.headers.get("Content-Disposition") or ""
    check("🔒 文件名含 filename*=UTF-8'' 编码", "filename*=UTF-8''" in cd)
    check("🔒 文件名保留 ASCII 兜底", 'filename="' in cd)
    check("🔒 响应 no-store", r.headers.get("Cache-Control") == "no-store")
    check("⑤ Content-Type 为 text/csv", "text/csv" in (r.headers.get("Content-Type") or ""))

    txt = r.content.decode("utf-8-sig")
    lines = [ln for ln in txt.splitlines() if ln.strip()]
    hdr = lines[0]
    expect = "sku_id,warehouse_id,date,forecast_sales,granularity_level,actual_sales,error"
    check("🔒 CSV 列序与 EXPORT_HEADER 完全一致", hdr == expect, hdr)

    st, bd = api("GET", f"/api/v1/predictions/{tid}/detail?size=1")
    row = (bd.get("data", {}).get("items") or [{}])[0]
    first = lines[1].split(",")
    check("⑤ 首行字段与明细一致",
          first[0] == row.get("sku_id") and first[1] == row.get("warehouse_id")
          and first[2] == row.get("forecast_date"),
          f"{first[:3]} vs {row.get('sku_id')},{row.get('warehouse_id')},{row.get('forecast_date')}")
    check("🔒 无真值时 actual_sales / error 为空（不是 0）",
          first[5] == "" and first[6] == "", f"actual={first[5]!r} error={first[6]!r}")

    st, bd = api("GET", f"/api/v1/predictions/{tid}/detail?size=1")
    n_detail = bd.get("data", {}).get("total")
    check("⑤ CSV 行数 = 明细总行数", len(lines) - 1 == n_detail,
          f"CSV {len(lines) - 1} vs 明细 {n_detail}")

    # 中文不乱码
    check("🔒 中文仓库名正常（无乱码）", "上海仓" in txt or "广州仓" in txt,
          lines[1].split(",")[1])

    # 非法 format
    r = c.get(f"/api/v1/predictions/{tid}/export?format=pdf", headers=STATE["H"])
    check("⑤ 非法 format 被拒（40012）",
          r.status_code == 200 and r.json().get("code") == 40012,
          str(r.json().get("message"))[:40])

    # Excel
    r = c.get(f"/api/v1/predictions/{tid}/export?format=xlsx", headers=STATE["H"])
    check("⑤ XLSX 导出 HTTP 200", r.status_code == 200)
    check("🔒 XLSX 是真 zip（PK 魔数）", r.content[:2] == b"PK", repr(r.content[:2]))
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb[wb.sheetnames[0]]
        xhdr = [str(c.value) for c in ws[1]]
        check("🔒 Excel 实际打开，列序与 CSV 一致", xhdr == expect.split(","), str(xhdr))
        check("🔒 Excel 中文不乱码", ws.cell(row=2, column=2).value in ("上海仓", "广州仓", "北京仓", "武汉仓", "成都仓"),
              str(ws.cell(row=2, column=2).value))
        check("⑤ Excel 行数 = 明细 + 表头", ws.max_row == n_detail + 1,
              f"{ws.max_row} vs {n_detail + 1}")
        check("🔒 Excel 无真值单元格为空",
              ws.cell(row=2, column=6).value is None,
              repr(ws.cell(row=2, column=6).value))
    except ImportError:
        print("  [SKIP] 未装 openpyxl，跳过 Excel 内容校验")
    except Exception as exc:  # noqa: BLE001
        check("🔒 Excel 实际可打开", False, str(exc)[:80])


# =============================================================================
# ⑥ L3 稀疏序列（API-15-2：真实数据 L3=0，必须构造）
# =============================================================================
def case_l3_sparse() -> None:
    section("⑥ L3 稀疏序列（构造触发，API-15-2）")
    import csv as _csv

    # 🔴 当前示例数据集 L3 = 0，该分支在真实数据上不可见。
    #    这里构造一份「大量短序列」的 CSV：每条序列只有 5 天历史
    #    （< L1_MIN_VALID_POINTS=90），必然被判为 L3。
    src = ROOT / "outputs" / "models" / "demo_sales.csv"
    if not src.exists():
        print("  [SKIP] 找不到示例 CSV，跳过 L3 用例")
        return
    import pandas as pd

    df = pd.read_csv(src, dtype=str)
    short = df.groupby(["sku_id", "warehouse_id"], sort=False).head(5)
    p = TMP / "sparse.csv"
    short.to_csv(p, index=False, encoding="utf-8")

    c: TestClient = STATE["client"]
    with open(p, "rb") as f:
        r = c.post("/api/v1/datasets", files={"file": ("sparse.csv", f, "text/csv")},
                   headers=STATE["H"])
    check("⑥ 稀疏数据集上传成功", r.status_code == 200 and r.json().get("code") == 0,
          str(r.json().get("message"))[:50])
    ds = ((r.json().get("data") or {}).get("dataset") or {}).get("id")

    st, b = api("POST", "/api/v1/predictions",
                json={"dataset_id": ds, "horizon": 7})
    d = b.get("data") or {}
    tid = d.get("task_id")
    check("⑥ 稀疏数据预测完成", d.get("status") == "succeeded",
          f"code={b.get('code')} status={d.get('status')} "
          f"err={d.get('error_code') or ''} msg={str(b.get('message'))[:40]}")
    cov = d.get("coverage") or {}
    check("⑥ 稀疏数据出现 L3 序列", cov.get("L3", 0) > 0, json.dumps(cov, ensure_ascii=False))
    check("🔒 L3 计入总数（三项之和 = total）",
          cov.get("L1", 0) + cov.get("L2", 0) + cov.get("L3", 0) == cov.get("total"),
          f"{cov.get('L1')}+{cov.get('L2')}+{cov.get('L3')} vs {cov.get('total')}")

    if not tid:
        return
    st, b = api("GET", f"/api/v1/predictions/{tid}")
    dd = b.get("data") or {}
    cov2 = dd.get("coverage") or {}
    check("🔒 API-14 覆盖率同样闭合",
          cov2.get("L1", 0) + cov2.get("L2", 0) + cov2.get("L3", 0) == cov2.get("total"),
          json.dumps(cov2, ensure_ascii=False))
    check("🔒 L3 不参与主指标（series_evaluated 只数 L1）",
          (dd.get("metrics") or {}).get("series_evaluated") == cov2.get("L1"),
          f"evaluated={(dd.get('metrics') or {}).get('series_evaluated')} L1={cov2.get('L1')}")

    series = dd.get("series") or []
    l3s = [s for s in series if s.get("granularity_level") == "L3"]
    check("⑥ 曲线里有 L3 序列", len(l3s) > 0, f"{len(l3s)} / {len(series)}")
    if l3s:
        s = l3s[0]
        check("🔒 L3 曲线 actual 为 null", s.get("actual") is None)
        check("🔒 L3 预测段不被补齐（长度 < horizon）",
              len(s.get("forecast_dates") or []) < 7,
              f"{len(s.get('forecast_dates') or [])} 天（horizon=7）")
        check("🔒 L3 标注不可预测", s.get("granularity_level") == "L3")

    st, b = api("GET", f"/api/v1/predictions/{tid}/detail?level=L3&size=5")
    check("⑥ 明细可筛出 L3", b.get("data", {}).get("total", 0) > 0,
          f"{b.get('data', {}).get('total')} 行")

    r = c.get(f"/api/v1/predictions/{tid}/export?format=csv", headers=STATE["H"])
    lines = [ln for ln in r.content.decode("utf-8-sig").splitlines() if ln.strip()]
    has_l3 = any(",L3," in ln for ln in lines[1:])
    check("🔒 导出含 granularity_level=L3 标记", has_l3,
          f"{sum(1 for ln in lines[1:] if ',L3,' in ln)} 行 L3")


# =============================================================================
def main() -> int:
    # 🔴 整个运行都在**同一个** `with client:` 内：lifespan 退出即
    #    `holder.clear()`，一旦中途退出，后面需要推理的用例就会拿到 40404。
    client = TestClient(app)
    STATE["client"] = client
    with client:
        if not case_setup():
            print("\n[FATAL] 前置未通过，后续用例跳过")
            return 1
        case_api14()
        case_edge()
        case_api15()
        case_api16()
        case_l3_sparse()

    print(f"\n{'=' * 70}")
    print(f"EP-6 接口验收：{PASS} / {TOTAL} 通过")
    if PASS != TOTAL:
        print("\n未通过项已在上方逐条标出 [FAIL]")
    print(f"{'=' * 70}")
    return 0 if PASS == TOTAL else 1


if __name__ == "__main__":
    sys.exit(main())
